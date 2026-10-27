import copy,hashlib,json,re,zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).parent
E=TMP/'evidence'
OUT=ROOT/'output/report'
REF=Path(r'C:/Users/ASUS/Downloads/BCSE302P_Mini_Project_Report_Template.docx')
data=json.loads((E/'report-data.json').read_text(encoding='utf8'))
proof=json.loads((E/'browser-booking.json').read_text(encoding='utf8'))
DATE='19 October 2026'
TITLE='QureFlow Clinic Queue and Appointment Management Using MySQL'
source=Document(REF)
doc=Document(REF)
cover_end=doc.paragraphs[21]._p
remove=False
for child in list(doc._element.body):
    if child is cover_end:remove=True
    if remove and child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
seen_styles={}
for style_element in list(doc.styles.element):
    identifier=style_element.get(qn('w:styleId'))
    if identifier:
        if identifier in seen_styles:doc.styles.element.remove(seen_styles[identifier])
        seen_styles[identifier]=style_element
for style in list(doc.styles):
    # The retained template duplicates IDs and uses noncanonical built-in names.
    if re.fullmatch(r'Heading [1-9]',style.name or ''):
        style.element.find(qn('w:name')).set(qn('w:val'),style.name.lower())
    if style.type==WD_STYLE_TYPE.PARAGRAPH:
        style.font.name='Times New Roman'
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(14 if style.name.startswith('Heading') or style.name=='Title' else 12)
        if style.name.startswith('Heading') or style.name=='Title':style.font.bold=True
if 'Normal' not in doc.styles:
    normal=doc.styles.add_style('Normal',WD_STYLE_TYPE.PARAGRAPH)
else:normal=doc.styles['Normal']
normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2']:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.bold=True;s.font.italic=False
    s.paragraph_format.keep_with_next=True;s.paragraph_format.line_spacing=1.5
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
    level=OxmlElement('w:outlineLvl');level.set(qn('w:val'),'0' if name=='Heading 1' else '1');s.element.get_or_add_pPr().append(level)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
if 'Front Heading' not in doc.styles:doc.styles.add_style('Front Heading',WD_STYLE_TYPE.PARAGRAPH)
front=doc.styles['Front Heading'];front.base_style=doc.styles['Heading 1'];front.paragraph_format.page_break_before=True
el=front.element.get_or_add_pPr();o=OxmlElement('w:outlineLvl');o.set(qn('w:val'),'9');el.append(o)
for s in doc.sections:
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)

def format_run(run,size=12,bold=False,italic=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.italic=italic;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
def p(text='',bold=False,italic=False,align=None,keep=False):
    x=doc.add_paragraph(style='Normal');x.paragraph_format.keep_with_next=keep
    x.paragraph_format.line_spacing=1.5
    x.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    format_run(x.add_run(text),bold=bold,italic=italic)
    return x
def h(text,level=1):
    x=doc.add_paragraph(text,style='Heading '+str(level));return x
def code(text):
    for line in text.strip().splitlines():
        x=p(line,align=WD_ALIGN_PARAGRAPH.LEFT)
        x.paragraph_format.space_after=Pt(0)
        x.paragraph_format.keep_together=True
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');x._p.get_or_add_pPr().append(sh)
    spacer=p('');spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3)
def caption(text):
    x=p(text,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
    x.paragraph_format.space_after=Pt(8)
    return x
def figure(name,text,width=6.2,maxheight=5.9):
    image=Image.open(E/name);w,h0=image.size
    width=min(width,maxheight*w/h0)
    x=p('',align=WD_ALIGN_PARAGRAPH.CENTER,keep=True)
    x.add_run().add_picture(str(E/name),width=Inches(width))
    x.paragraph_format.space_after=Pt(0)
    pic=doc.inline_shapes[-1]._inline.docPr;pic.set('descr',text)
    caption(text)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths is None:widths=[6.26/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=str(value)
    for row in rows:
        for c,value in zip(t.add_row().cells,row):c.text=str(value)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
        if ri==0:
            rep=OxmlElement('w:tblHeader');pr.append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'60');margin.set(qn('w:type'),'dxa');margins.append(margin)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000');borders.append(b)
            tcpr.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'D9E2F3');tcpr.append(sh)
            for para in cell.paragraphs:
                para.style=normal;para.paragraph_format.line_spacing=1.5;para.paragraph_format.space_after=Pt(0)
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:format_run(run,bold=ri==0)
    spacer=p('');spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3)
    return t

# Edit retained cover slots in place.
replacements={0:'VELLORE INSTITUTE OF TECHNOLOGY VELLORE',1:'School of Computer Science and Engineering (SCOPE)',
4:'MINI PROJECT REPORT',6:TITLE,8:'BCSE302P Database Systems Laboratory',11:'Submitted by',
15:'Under the guidance of',16:'Dr. G. M. Karthik',
17:'Associate Professor, School of Computer Science and Engineering (SCOPE)',
18:'Vellore Institute of Technology, Vellore',20:'Date of Submission: '+DATE}
cover=list(doc.paragraphs)
for i,text in replacements.items():
    para=cover[i];para.clear();para.style=doc.styles['Title'] if i==6 else normal
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1.5
    para.paragraph_format.space_after=Pt(4)
    format_run(para.add_run(text),size=14 if i in [0,4,6] else 12,bold=i in [0,4,6,16,20],italic=i in [11,15])
cover[13]._element.getparent().remove(cover[13]._element)
members=doc.tables[0]
for row in list(members.rows)[3:]:members._tbl.remove(row._tr)
for row,values in zip(list(members.rows)[1:],[['1','Kartikay Gupta','24BCI0109'],['2','Sreyanko Sinha','24BCB0059']]):
    for cell,val in zip(row.cells,values):cell.text=val
for ri,row in enumerate(members.rows):
    for c in row.cells:
        for para in c.paragraphs:
            para.paragraph_format.line_spacing=1.5
            for r in para.runs:format_run(r,bold=ri==0)

h('CERTIFICATE')
p('This is to certify that the Mini Project Report titled '+TITLE+' submitted by Kartikay Gupta (24BCI0109) and Sreyanko Sinha (24BCB0059), in partial fulfilment of the requirements of BCSE302P Database Systems Laboratory, is a record of their project work during the Fall Semester 2026-27.')
p('Declaration of originality',bold=True)
p('We declare that this report describes our QureFlow project and its locally executed MySQL demonstration. The test identities and clinical organisations are synthetic. External technical documentation is acknowledged in the references, and execution results are distinguished from features requiring further device verification.')
p('Kartikay Gupta  |  24BCI0109')
p('Sreyanko Sinha  |  24BCB0059')
p('')
p('Signature of the Faculty',bold=True,align=WD_ALIGN_PARAGRAPH.RIGHT)
p('Dr. G. M. Karthik',align=WD_ALIGN_PARAGRAPH.RIGHT)
p('Associate Professor, SCOPE',align=WD_ALIGN_PARAGRAPH.RIGHT)
p('VIT, Vellore',align=WD_ALIGN_PARAGRAPH.RIGHT)
p('Place: Vellore')
p('Date: '+DATE)

h('ABSTRACT')
abstract=('QureFlow is a clinic queue and appointment management application designed to make visit planning and queue progression more transparent. This project presents its MySQL database variant for the Database Systems Laboratory. The implementation separates clinics, doctors, specialties, languages, users, appointments and queue tokens into relational tables with explicit keys and constraints. SQL joins, filters, ordering and pagination support doctor selection, while InnoDB transactions and row locks protect appointment capacity and concurrent queue operations. A stored function measures consultation duration, a procedure produces daily patient-flow results, three triggers enforce and record queue transitions, and an explicit cursor returns doctor workload. A React web interface communicates with an Express API through parameterised database queries. The same interface is packaged as an Android application, with local server discovery and anonymous connection verification. Testing covers valid workflows, invalid transitions, duplicate visits, clinic isolation, injection inputs and simultaneous bookings. All 24 MySQL integration cases passed in the recorded run, and browser booking and cancellation were verified against SQL records. The result is a reproducible database demonstration; production deployment and final hotspot acceptance on Android remain future verification tasks.')
assert 150<=len(abstract.split())<=200,len(abstract.split())
p(abstract)
p('Keywords: MySQL, queue management, normalization, transactions, triggers, Android',italic=True)

doc.add_paragraph('TABLE OF CONTENTS',style='Front Heading')
toc=p('')
for kind in ['begin','instruction','separate','end']:
    run=OxmlElement('w:r')
    if kind=='instruction':
        e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' TOC \\o "1-2" \\h \\z \\u '
    elif kind=='separate':
        e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'separate')
    else:
        e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),kind)
    run.append(e);toc._p.append(run)

h('1 INTRODUCTION')
h('1.1 Problem Statement',2)
p('A manual clinic queue gives patients little visibility into their position and gives reception staff limited protection against duplicate entries or conflicting calls. Appointment records, doctor schedules and queue state can also diverge when they are maintained in separate spreadsheets. QureFlow brings these records into one database-backed workflow so that a patient can select a doctor and book a visit, while staff manage the corresponding queue.')
p('The objective is to reduce avoidable uncertainty and waiting through better scheduling and queue visibility. The recorded experiments verify data integrity and workflow behaviour; they do not constitute a clinical study or quantify a reduction in real hospital waiting time.')
h('1.2 Objectives',2)
for n,text in enumerate([
'Design relational master data for clinics, doctors, specialties, languages and users.',
'Perform doctor selection, filtering, ordering and pagination through SQL queries.',
'Create appointments and tokens atomically while protecting slot capacity under concurrent requests.',
'Enforce queue transitions and audit changes with database constraints and triggers.',
'Integrate patient and staff web screens with the MySQL API and package the interface for Android.',
'Verify database objects and negative cases through repeatable tests and captured execution output.'],1):p(str(n)+'. '+text)
h('1.3 Importance of DBMS',2)
p('Primary and foreign keys connect records without storing the whole application as document blobs. Unique indexes prevent duplicate identities and conflicting active visits. Transactions let a booking either create all linked records or roll back. Row locks serialize competing updates to a doctor or appointment slot. Parameterised queries keep input values separate from SQL syntax, and application authorization restricts records to the permitted user and clinic. These controls address different failure modes and are used together.')
h('1.4 Scope of the Project',2)
p('The scope includes a local MySQL 8.4 database, patient and staff authentication for synthetic accounts, a doctor directory, scheduled appointments, queues, cancellation, audit records, daily analytics and an Android debug package. The earlier AWS DynamoDB/Cognito variant is retained separately; it is not the database or authentication path evaluated in this report.')
p('Clinical diagnosis, automatic medical triage, billing, electronic health-record interoperability and production hospital deployment are outside this laboratory submission. Consultation estimates are operational guidance. Actual Android hotspot discovery remains a device acceptance item.')

h('2 SYSTEM REQUIREMENTS')
h('2.1 Hardware Requirements',2)
p('The local execution workstation uses Windows 11 and 16 GB RAM. The following is a practical development recommendation, not a measured minimum specification.')
table(['Component','Specification'],[
['Processor','64-bit processor with virtualization enabled for Docker Desktop'],
['Memory','8 GB or more recommended; 16 GB used for the recorded run'],
['Storage','At least 10 GB free for dependencies, database volume and build tools'],
['Network','Local Wi-Fi or phone hotspot for Android-to-computer testing'],
['Optional device','Android API 24 or newer for installing the debug APK']],[1.65,4.61])
h('2.2 Software Requirements',2)
table(['Category','Tool or version'],[
['Database','MySQL '+data['version']+' with InnoDB, in the mysql:8.4 Docker image'],
['Database client','MySQL command-line client inside the container; Workbench is optional'],
['Backend','Node.js, TypeScript, Express 4 and mysql2/promise'],
['Web interface','React 18, TypeScript 5, Vite 6 and Tailwind CSS'],
['Authentication','Local synthetic accounts, bcrypt password hashes and JWT sessions'],
['Android','Capacitor 8.5.3, SDK API 36, Gradle 8.14.3 and JDK 22'],
['Verification','Node test runner, TypeScript execution and browser automation'],
['Local endpoints','Web: localhost:5174; API: port 5001; MySQL: 127.0.0.1:3307']],[1.65,4.61])

h('3 DATABASE DESIGN')
h('3.1 Entity Relationship Diagrams',2)
p('The database contains 18 base tables and four views. The diagrams divide the model into catalog data, visit data and operational state to keep names and cardinalities readable. Arrows follow parent-to-child relationships. A 1:N label permits zero or more children for a parent; a non-null foreign key requires each child to reference one parent.')
figure('er-catalog.png','Figure 3.1 Catalog relations and doctor membership tables',maxheight=5.5)
p('Every doctor belongs to one clinic and one specialty. Languages form a many-to-many relationship through doctor_languages. The doctor_days and doctor_interests tables remove repeating lists. A clinic has at most one settings row.')
figure('er-visits.png','Figure 3.2 Appointment and token relationships',maxheight=5.8)
p('Appointments and tokens reference users and doctors. The patient role is checked by the API; the foreign key itself references users. A nullable unique appointmentId allows a token to be unlinked and prevents several tokens from claiming one appointment. The dashed slot match is maintained by application transactions and triggers; it is not a declared composite foreign key.')
figure('er-state.png','Figure 3.3 Per-doctor state and stored operational models',maxheight=3.8)
p('Counters, pause state and prediction models reference doctors. The dashed clinic-to-model line denotes a scope value without a declared foreign key. Optional staff assignments in users reference a clinic and doctor. Audit events store correlated identifiers without foreign keys; schema_migrations and demo_seed are independent support tables. All relations are listed below.')
h('3.2 Relational Schema',2)
p('Tables 3.1 to 3.18 list key and selected operational attributes for every base relation. The complete DDL in Appendix A defines all 199 columns, enumerated values, defaults, generated expressions and constraints. PK means primary key, FK means foreign key and UQ means a unique constraint or index.')
selected={
'clinics':['id','name','district','phone'],
'specialties':['id','name'],'languages':['id','name'],
'doctors':['id','clinicId','specialtyId','name','experienceYears','consultationFee'],
'doctor_days':['doctorId','dayOfWeek'],'doctor_languages':['doctorId','languageId'],'doctor_interests':['doctorId','interest'],
'users':['id','name','phone','email','role','clinicId','doctorId'],
'clinic_settings':['clinicId','averageConsultationMinutes','bufferMinutes','allowWalkIns'],
'queue_counters':['doctorId','queueDate','currentNumber'],'queue_state':['doctorId','isPaused'],
'appointment_slots':['doctorId','slotDate','slotMinute','bookedCount','capacity'],
'appointments':['id','patientId','doctorId','date','slotMinute','status'],
'queue_tokens':['id','appointmentId','patientId','doctorId','queueDate','tokenSeq','status','activePatientId','servingDoctorId'],
'audit_events':['id','clinicId','entityType','entityId','action','actorId','requestId'],
'prediction_models':['doctorId','clinicId','model','updatedAt'],
'schema_migrations':['version','appliedAt'],'demo_seed':['version','seededAt']}
descriptions={
'id':'Record identifier','name':'Display name','clinicId':'Clinic reference or audit scope','specialtyId':'Specialty reference',
'patientId':'Patient user reference','doctorId':'Doctor reference','appointmentId':'Optional linked appointment',
'phone':'Contact number','email':'Unique contact when present','role':'Patient or permitted staff role',
'experienceYears':'Experience in years, 0 to 70','consultationFee':'Non-negative consultation fee',
'dayOfWeek':'Consultation weekday','languageId':'Language reference','interest':'Doctor care interest',
'averageConsultationMinutes':'Configured duration baseline','bufferMinutes':'Non-negative slot buffer','allowWalkIns':'Walk-in configuration',
'queueDate':'Clinic-local queue date','currentNumber':'Last allocated sequence','isPaused':'Queue pause flag',
'slotDate':'Appointment date','slotMinute':'Minutes after midnight','bookedCount':'Current reserved places','capacity':'Maximum slot places',
'date':'Scheduled appointment date','status':'Allowed visit state','tokenSeq':'Per-doctor daily sequence',
'activePatientId':'Generated active-patient claim','servingDoctorId':'Generated current-consultation claim',
'entityType':'Audited entity category','entityId':'Audited record identifier','action':'Recorded operation',
'actorId':'Initiating user or SQL client','requestId':'Correlation identifier','model':'Operational model JSON',
'updatedAt':'Last modification timestamp','version':'Migration or fixture version','appliedAt':'Migration application time',
'seededAt':'Fixture installation time','district':'Clinic district'}
fks={(r['TABLE_NAME'],r['COLUMN_NAME']):r['REFERENCED_TABLE_NAME'] for r in data['foreignKeys']}
unique={(r['TABLE_NAME'],r['columns']) for r in data['indexes'] if not r['NON_UNIQUE'] and ',' not in r['columns']}
for num,(name,names) in enumerate(selected.items(),1):
    p('Table 3.'+str(num)+' '+name,bold=True,keep=True,align=WD_ALIGN_PARAGRAPH.LEFT)
    rows=[]
    cols={c['COLUMN_NAME']:c for c in data['columns'] if c['TABLE_NAME']==name}
    for key in names:
        c=cols[key];typ=c['COLUMN_TYPE']
        if typ.startswith('enum('):typ='ENUM'
        flags=[]
        if c['COLUMN_KEY']=='PRI':flags.append('PK')
        elif (name,key) in unique:flags.append('UQ')
        if c['IS_NULLABLE']=='NO':flags.append('NOT NULL')
        if (name,key) in fks:flags.append('FK '+fks[(name,key)])
        if any(value in c['EXTRA'] for value in ['STORED GENERATED','VIRTUAL GENERATED']):flags.append('Generated')
        display_key=re.sub(r'(?<=[a-z])(?=[A-Z])','\u200b',key)
        rows.append([display_key,typ,', '.join(flags) or 'Nullable',descriptions.get(key,'Operational attribute')])
    table(['Attribute','Data type','Constraint','Description'],rows,[2.05,1.05,1.45,1.71])
h('3.3 Schema Definition and Constraints',2)
p('The schema is created by versioned SQL files at startup. Important rules include unique non-null contact values, non-negative fees, valid schedule bounds, a bookedCount within capacity and foreign keys that reject orphan visits. MySQL uses InnoDB for these relations [1].')
ddl=(ROOT/'backend/sql/001_schema.sql').read_text(encoding='utf8')
slot=re.search(r'CREATE TABLE IF NOT EXISTS appointment_slots \([\s\S]+?\n\);',ddl).group()
code(slot)
p('The generated activePatientId is populated only for waiting/current tokens; a composite unique index with queueDate enforces one active visit per patient per day. Similarly, servingDoctorId is populated only for current tokens. A unique index on servingDoctorId and queueDate permits at most one current consultation per doctor per day. Token numbering is unique on doctorId, queueDate and tokenSeq.')
code("UNIQUE active_visit (activePatientId, queueDate),\nUNIQUE one_consultation (servingDoctorId, queueDate),\nUNIQUE token_sequence (doctorId, queueDate, tokenSeq)")
h('3.4 Normalization',2)
p('An unnormalized visit sheet could repeat doctor and clinic details beside each patient, while keeping language and weekday lists in single cells. Updating a clinic address or deleting the last visit could then change or remove unrelated master information. The catalog decomposition removes these dependencies.')
table(['Normal form','Relations and dependencies','Anomaly addressed'],[
['1NF','Atomic master attributes; doctor_days and doctor_languages hold one membership per row.','Repeating weekday and language groups are removed.'],
['2NF','In doctor_languages, the key is (doctorId, languageId). Doctor details stay in doctors; language names stay in languages.','Non-key attributes do not depend on only part of the membership key.'],
['3NF','doctorId determines clinicId and specialtyId; their descriptive attributes are stored in clinics and specialties.','Doctor rows do not repeat clinic address or specialty name.'],
['BCNF assessment','Simple lookup and membership relations satisfy the stated key dependencies. No blanket BCNF claim is made for the operational schema.','Operational copies and derived state are documented rather than hidden.']],[1.0,2.70,2.56])
p('Key dependencies include doctorId -> doctor attributes, clinicId -> clinic attributes and (doctorId, slotDate, slotMinute) -> capacity, bookedCount. Doctor, clinic and specialty details are stored separately. Nullable email and phone are unique only when present.')
p('Operational duplication remains: tokens repeat appointment details, clinic_settings copies a clinic name, and occupancy/counters are derived. Transactions and triggers maintain tested updates. Direct SQL edits require the same consistency rules.')

h('4 IMPLEMENTATION')
p('SQL source is maintained in backend/sql. The following snippets are taken from the project implementation. The function and daily-flow procedure are used by analytics; triggers execute as part of queue writes. The cursor is a read-only laboratory demonstration. The DBMS is MySQL, so these are MySQL stored programs rather than Oracle PL/SQL.')
routines=(ROOT/'backend/sql/002_routines.sql').read_text(encoding='utf8')
def routine(kind,name):
    return 'DELIMITER $$\n'+re.search(r'CREATE '+kind+r' '+re.escape(name)+r'[\s\S]+?END\$\$',routines).group()+'\nDELIMITER ;'
h('4.1 Functions',2)
p('fn_consultation_minutes accepts two DATETIME(3) values and returns a DECIMAL(8,2) duration. It returns NULL for missing, reversed or longer-than-two-hour intervals and returns zero for an identical timestamp. The analytics query calls it when aggregating completed visits. This gives reporting one shared duration rule [2].')
code(routine('FUNCTION','fn_consultation_minutes'))
p('The executed positive case returned 12.50 minutes. Missing and reversed inputs returned NULL.')
figure('evidence-function.png','Figure 4.1 MySQL terminal execution of the consultation-duration function',maxheight=2.8)
h('4.2 Procedures',2)
p('sp_daily_patient_flow takes a clinic identifier and an inclusive date range. A recursive date series is left-joined with daily_patient_flow, so a day without matching visits is returned with zero counts instead of disappearing from the chart. The backend calls this procedure for its 14-day analytics series [2].')
code(routine('PROCEDURE','sp_daily_patient_flow'))
figure('evidence-procedure.png','Figure 4.2 MySQL terminal procedure output including zero-volume days',maxheight=5.6)
h('4.3 Packages',2)
p('MySQL 8.4 does not implement Oracle-style package specifications and bodies. Consequently, this project does not claim a PL/SQL package. The report retains this template section to make the platform choice explicit. Related SQL routines are deployed together in 002_routines.sql, while connection and repository functions are grouped in TypeScript modules; those modules are not database packages.')
code("SELECT ROUTINE_NAME, ROUTINE_TYPE\nFROM information_schema.ROUTINES\nWHERE ROUTINE_SCHEMA = DATABASE();")
p('The recorded database contains one stored function and two stored procedures. Package syntax from the Oracle examples in the template has not been inserted into the MySQL project.')
h('4.4 Triggers',2)
p('All three triggers operate at row level on queue_tokens. validate_token_transition runs BEFORE UPDATE and rejects invalid state changes or calling a patient while the doctor queue is paused. audit_token_insert runs AFTER INSERT and creates an audit event. synchronize_token_update runs AFTER UPDATE, maps token state to appointment state, releases cancelled/skipped slot occupancy and records the transition [3].')
code(routine('TRIGGER','validate_token_transition'))
code(routine('TRIGGER','synchronize_token_update'))
p('The trigger demonstration uses disposable synthetic records inside one transaction. Before cancellation the appointment was booked, the token waiting and bookedCount equal to one. After the token update, both records were cancelled and bookedCount was zero. The final ROLLBACK left zero demonstration users, appointments, tokens and audit events. The browser workflow independently verified the same linked-state behaviour on token '+proof['appointment']['tokenNumber']+'.')
h('4.5 Cursors',2)
p('sp_doctor_workload_cursor declares an explicit cursor over the clinic roster. It opens the cursor, fetches one doctor at a time, counts waiting tokens for the requested date and appends a JSON result. A NOT FOUND handler terminates the loop, and the cursor is closed before the result is returned. It demonstrates procedural row traversal; normal operational summaries use set-based aggregation [4].')
code(routine('PROCEDURE','sp_doctor_workload_cursor'))
figure('evidence-cursor.png','Figure 4.3 Cursor procedure and its JSON result in the MySQL terminal',maxheight=4.6)
h('4.6 SQL Search and Transaction Control',2)
p('Directory search is performed by MySQL, including specialty and language membership, clinic scope, ordering and LIMIT/OFFSET. Values are bound as parameters. Sort expressions are selected from a whitelist rather than concatenated from free-form user input.')
code("SELECT d.id, d.name, s.name AS specialty,\n       c.name AS clinic, d.consultationFee\nFROM doctors d\nJOIN specialties s ON s.id = d.specialtyId\nJOIN clinics c ON c.id = d.clinicId\nWHERE s.name = 'Gynecology' AND d.status = 'active'\nORDER BY d.consultationFee, d.name;")
figure('evidence-directory.png','Figure 4.4 Terminal SQL join, specialty filter and fee-sorted result',maxheight=3.8)
p('Booking uses a pooled connection with BEGIN, COMMIT and ROLLBACK. The transaction locks the doctor and slot rows, checks capacity, inserts linked appointment/token records and updates occupancy. Queue calls and pause operations use a consistent doctor-first lock order. This avoids implementing concurrency protection solely as a JavaScript pre-check [5].')
code("SELECT status FROM doctors WHERE id = ? FOR UPDATE;\nSELECT bookedCount, capacity\nFROM appointment_slots\nWHERE doctorId = ? AND slotDate = ? AND slotMinute = ?\nFOR UPDATE;")

h('5 FRONT END AND ANDROID INTEGRATION')
h('5.1 User Interface',2)
p('Patient screens expose the directory, available slots, booking confirmation, appointment history and queue information. Staff screens expose clinic-scoped doctors, appointments, queues and analytics. The following images are actual screenshots of the running local web application, captured after API responses completed.')
figure('web-specialty.png','Figure 5.1 Web directory filtered to Gynecology using database results',maxheight=4.8)
p('The specialty selection is sent to the SQL search endpoint. The displayed doctors come from joined clinic and specialty records. Filtering and the selected sort order are applied by MySQL, with paginated results returned to the browser.')
figure('web-booking-confirmation.png','Figure 5.2 Patient confirmation before submitting the verification booking',maxheight=4.8)
p('The confirmation page shows the selected doctor, date and slot before any write is submitted. For this execution, the signed-in synthetic patient selected an 08:00 appointment on 20 October. The browser then submitted one booking request; the returned identifier was used for the independent SQL check.')
figure('web-booking-persisted.png','Figure 5.3 Confirmed browser booking for 20 October 2026',maxheight=4.8)
p('The appointment history displays the returned token A-001. At this point SQL showed an appointment in booked state, a linked waiting token and one reserved place. The comparison in Section 6 records the exact before and after values.')
figure('web-staff-doctors.png','Figure 5.4 Clinic-scoped doctor administration after staff sign-in',maxheight=4.8)
p('Staff sign-in exposes the assigned clinic workflow. SQL queries and API authorization jointly apply clinic scope. Cross-clinic modification and patient attempts to use staff-only appointment filters were rejected in the recorded regression cases.')
figure('web-analytics.png','Figure 5.5 Clinic analytics summary computed from recorded SQL data',maxheight=4.8)
h('5.2 Integration with the DBMS',2)
p('The React client sends HTTP requests to Express. The API validates the session and role, applies clinic scope and calls the MySQL repository through mysql2/promise. Database triggers run within the same transaction as the write that caused them. The API returns success only after the database operation has completed. Database passwords remain in the backend configuration.')
code("const [rows] = await connection.execute(\n  'SELECT version FROM users WHERE id=? FOR UPDATE',\n  [userId]\n);")
p('The connection pool is configured for the local MySQL endpoint, uses UTC timestamps and disables multiple SQL statements per request. On a failed transaction the connection rolls back, resets the audit context variables and returns to the pool. The database stores clinic-local dates separately from event timestamps.')
h('5.3 Android Application',2)
p('Capacitor packages the React interface as com.qureflow.dbms. The verified debug APK is version 1.3, code 4, targeting SDK 36 with minimum API 24. It calls the same API rather than connecting directly to MySQL [6].')
p('Connection discovery combines mDNS with a bounded UDP fallback for Wi-Fi/hotspot interfaces. Candidate addresses receive anonymous health checks before a saved clinic is selected. Reconnection pauses requests and cancels in-flight calls; interrupted writes are not replayed. Doctor reads reload after verification without resetting filters.')
figure('responsive-web.png','Figure 5.6 Responsive browser view at phone dimensions; this is not a native APK screenshot',width=2.3,maxheight=4.3)
p('APK compilation, package contents, signing and the absence of SQL credentials were verified. Hotspot discovery failed in the earlier build; final acceptance of the new Android hotspot path remains pending. A responsive browser image does not replace a native device execution test [7].')

h('6 VERIFICATION AND TESTING')
h('6.1 Test Cases',2)
p('The MySQL suite creates a UUID-named test database, loads isolated synthetic records, executes the API/database assertions and removes the test database. It does not run these destructive test cases against the normal demonstration schema. The fresh run completed 24 individual cases with no failures; Node also reports the parent suite as one passing test entry.')
cases=[
('Joined directory','Specialty, Hindi and experience order','Correct filtering and order','Matched expected rows'),
('SQL injection','Injected search and sort expression','No injected results; invalid sort rejected','Search empty; HTTP 400'),
('Pagination','Unknown specialty and page 999','Empty page with accurate total','Totals preserved'),
('Authorization','Anonymous read and cross-clinic edits','401 or 403 rejection','Requests denied'),
('Duplicate visit','Two concurrent joins for one patient','One active visit only','One succeeded'),
('Foreign key','Missing patient reference','Reject orphan token','FK error observed'),
('Priority order','Normal and emergency entries','Explicit emergency first','Correct next token'),
('Consultation claim','Second current token for doctor/date','Unique constraint rejects second','One current token'),
('Patient privacy','Another patient queue/cancel request','Conceal identity and deny mutation','Identity hidden; 403'),
('Trigger transition','Complete then reopen token','Audit and invalid-state rejection','Reopening rejected'),
('Pause and resume','Call paused queue, then resume','Reject while paused; succeed later','Both assertions passed'),
('Atomic booking','Book available future slot','Linked appointment/token; occupancy 1','Linked rows verified'),
('Full capacity','Book after capacity reached','Conflict without orphan records','Transaction rolled back'),
('Cancellation','Cancel linked visit','Both cancelled; occupancy released','Both states; count 0'),
('Profile concurrency','Stale version and duplicate email','Reject conflicting updates','Both rejected'),
('Identity protection','New ID with existing identity','Existing account not overwritten','Original unchanged'),
('Staff pagination','Clinic-scoped appointment page','Correct page; patient forbidden','Rows and 403 verified'),
('Stored objects','Function, procedure, views, cursor','12.5 minutes; 14 days; roster','Expected outputs matched'),
('Scoped filters','Injected patient filter and clinic read','No escape from authorized scope','Empty/authorized rows'),
('Buffer capacity','Change buffer and reuse slot','Newly opened place can be booked','Booking succeeded'),
('CHECK constraints','Negative fee and excess occupancy','Reject invalid rows','CHECK errors observed'),
('Invalid timing','Missing, reversed and long intervals','NULL results; zero for equal times','Expected values matched'),
('Concurrent capacity','Three patients, one remaining place','Exactly one successful booking','One success; count 1'),
('API persistence','HTTP booking then cancellation','SQL linked rows and both states','201 then 200; SQL matched')]
assert len(cases)==len(data['testCases'])==24
for start in [0,8,16]:
    p('Table 6.'+str(start//8+1)+' Executed database cases '+str(start+1)+' to '+str(start+8),bold=True,keep=True)
    table(['TC ID','Description','Input','Expected output','Actual output','Result'],[
      ['TC'+str(i+1).zfill(2),*cases[i],'Pass'] for i in range(start,start+8)
    ],[.48,1.08,1.32,1.39,1.34,.65])
h('6.2 Output Screenshots',2)
p('The following terminal captures show the executed test summary, SQL commands with their returned database values, and the saved cancellation state. The MySQL client prints each SQL statement before its result table.')
figure('evidence-tests.png','Figure 6.1 Terminal view of the executed test log showing zero failures',maxheight=4.8)
figure('evidence-schema.png','Figure 6.2 SQL commands and results for the database version and record counts',maxheight=4.6)
figure('evidence-booking.png','Figure 6.3 Terminal SQL join verifying the saved browser cancellation',maxheight=4.6)
figure('web-booking-cancelled.png','Figure 6.4 Browser appointment history after the verification cancellation',maxheight=4.8)
h('6.3 Validation of Functionality',2)
p('The browser confirmation produced appointment '+proof['appointment']['id']+' and token '+proof['appointment']['tokenNumber']+'. The independent SQL read found appointmentStatus=booked, tokenStatus=waiting and bookedCount=1. Cancellation through the browser changed both statuses to cancelled and released the slot to bookedCount=0. The final database contains 6,196 appointments and 6,196 queue tokens, including this cancelled verification record.')
p('The SQL function returned 12.50 minutes; invalid timing inputs produced NULL. The daily procedure returned all 14 requested dates, including zero-volume days. The cursor returned the eight doctors assigned to clinic-1 for the selected fixture date. Separate negative tests verified foreign keys, CHECK constraints, illegal transitions, identity conflicts and capacity races.')
p('Additional recorded checks passed 18 mobile/client tests and five discovery tests. These validate reconnection logic, protocol behaviour and local discovery exchange. They do not prove delivery of multicast or broadcast packets on every Android hotspot. APK package verification is kept separate from device acceptance.')

h('7 RESULTS AND DISCUSSION')
h('7.1 Achievements',2)
table(['Recorded item','Result'],[
['Database model','18 base tables, 199 columns, 18 declared foreign keys and four views'],
['Demonstration scale','Six fictional clinics, 48 doctors and 737 patient profiles'],
['Visit records','6,196 appointments and 6,196 tokens after browser verification'],
['Stored objects','One function, two procedures and three triggers; one procedure uses a cursor'],
['SQL verification','All 24 individual MySQL integration cases passed'],
['Web integration','Browser booking and cancellation matched independent SQL reads'],
['Android packaging','Version 1.3 built and signed; device hotspot acceptance pending']],[1.65,4.61])
p('The evidence demonstrates relational consistency, SQL filtering and rejection of conflicting writes. It supports a reproducible local workflow, without establishing a real-world waiting-time reduction or production certification.')
h('7.2 Challenges Faced',2)
p('Migrating the workflow required explicit relations and SQL queries in place of document-oriented access patterns. Concurrent booking needed row locks and conditional constraints; cancellation needed one consistent update to the token, appointment and capacity counter. Failed operations were verified to roll back.')
p('The phone could reach the API while hotspot discovery failed. A direct discovery fallback was added and locally tested. A cancelled doctor search also needed to reload after reconnection; a component regression verified the read refresh.')
h('7.3 Possible Improvements',2)
p('Next steps are native hotspot acceptance, HTTPS enrollment, backup and load tests, restricted database privileges and reconciliation of direct edits.')

h('8 CONCLUSION')
h('8.1 Summary of Learning Outcomes',2)
p('QureFlow applies relational design to a queue-management workflow. The project combines normalized master data with explicit operational state, uses SQL joins and grouping for useful application queries, and protects writes with transactions, locks, foreign keys, checks, unique indexes and triggers. It also distinguishes a reusable stored function, a set-based reporting procedure and an explicit cursor procedure.')
p('The most important outcome is reproducible verification. Positive and negative database cases, a real browser booking/cancellation and the matching SQL output demonstrate how interface actions affect stored records. The work also shows why API authorization, schema rules and network verification address different parts of the application rather than replacing one another.')
h('8.2 Future Scope',2)
p('The local demonstration can be extended toward clinic deployment after device, operational and security acceptance. Waiting-time models should be evaluated on consented operational timings and used as transparent scheduling guidance. They should not infer diagnoses or perform clinical triage. The original AWS variant and this MySQL laboratory variant should continue to be maintained as clearly separated configurations.')

h('9 REFERENCES')
refs=[
'[1] Oracle, "FOREIGN KEY Constraints," MySQL 8.4 Reference Manual. [Online]. Available: https://dev.mysql.com/doc/refman/8.4/en/create-table-foreign-keys.html. Accessed: Oct. 19, 2026.',
'[2] Oracle, "CREATE PROCEDURE and CREATE FUNCTION Statements," MySQL 8.4 Reference Manual. [Online]. Available: https://dev.mysql.com/doc/refman/8.4/en/create-procedure.html. Accessed: Oct. 19, 2026.',
'[3] Oracle, "CREATE TRIGGER Statement," MySQL 8.4 Reference Manual. [Online]. Available: https://dev.mysql.com/doc/refman/8.4/en/create-trigger.html. Accessed: Oct. 19, 2026.',
'[4] Oracle, "Cursors," MySQL 8.4 Reference Manual. [Online]. Available: https://dev.mysql.com/doc/refman/8.4/en/cursors.html. Accessed: Oct. 19, 2026.',
'[5] Oracle, "Locking Reads," MySQL 8.4 Reference Manual. [Online]. Available: https://dev.mysql.com/doc/refman/8.4/en/innodb-locking-reads.html. Accessed: Oct. 19, 2026.',
'[6] Capacitor, "Custom Native Android Code," Capacitor Documentation, v8. [Online]. Available: https://capacitorjs.com/docs/android/custom-code. Accessed: Oct. 19, 2026.',
'[7] Google, "Use network service discovery," Android Developers. [Online]. Available: https://developer.android.com/develop/connectivity/wifi/use-nsd. Accessed: Oct. 19, 2026.',
'[8] K. Gupta and S. Sinha, "QureFlow MySQL DBMS variant," project source and local execution records, 2026. Internal files: backend/sql/001_schema.sql, 002_routines.sql, backend/src/repositories/mysqlRepo.ts and backend/tests/mysql.test.ts.'
]
for text in refs:p(text,align=WD_ALIGN_PARAGRAPH.LEFT)

h('10 APPENDICES')
h('Appendix A Complete Schema Definition',2)
p('The following is the executed project schema source. It defines all base relations and the four views; the report dictionary in Section 3.2 lists selected attributes for readability.')
code(ddl)
h('Appendix B Additional Code and Source Map',2)
p('The transactional trigger demonstration is maintained in 004_transaction_demonstration.sql. Its disposable records are rolled back. The following excerpt is the actual cancellation and verification step.')
transaction=(ROOT/'backend/sql/004_transaction_demonstration.sql').read_text(encoding='utf8')
code(transaction[transaction.index("UPDATE queue_tokens SET status='cancelled'"):])
table(['Source file','Purpose'],[
['backend/sql/001_schema.sql','Relations, keys, checks, indexes and views'],
['backend/sql/002_routines.sql','Function, two procedures and three triggers'],
['backend/sql/003_demonstration_queries.sql','Manual search, aggregation and stored-object calls'],
['backend/sql/004_transaction_demonstration.sql','Trigger synchronization and rollback verification'],
['backend/src/dbms/connection.ts','Connection pool, transaction boundaries and error mapping'],
['backend/src/repositories/mysqlRepo.ts','Parameterized queries, locking and persistence'],
['backend/tests/mysql.test.ts','24 isolated SQL/API regression cases'],
['src/pages/patient/DoctorList.tsx','SQL-backed directory filters and reconnect refresh'],
['android/app/src/main/java/com/qureflow/dbms','Native navigation and discovery bridge']],[2.78,3.48])
h('Appendix C Terminal Demonstration',2)
p('With Docker Desktop and the MySQL container running, open a separate PowerShell window. This command can be run from any folder and prompts for the local demonstration database password.')
code('docker exec -it qureflow-dbms-mysql-1 mysql -uqureflow -p qureflow_dbms')
p('At the mysql prompt, execute each statement with a terminating semicolon. The client uses the database directly; no patient endpoint accepts arbitrary SQL.')
code("SHOW TABLES;\nDESCRIBE doctors;\nSELECT COUNT(*) AS totalDoctors FROM doctors;\nSELECT name,experienceYears,consultationFee\nFROM doctors ORDER BY consultationFee LIMIT 10;\nSELECT fn_consultation_minutes(\n  '2026-10-19 10:00:00','2026-10-19 10:12:30'\n) AS minutes;\nSHOW TRIGGERS;\nexit;")
h('Appendix D Evidence Provenance',2)
p('All web screenshots were captured from the running local application. SQL figures are native Windows console captures of the MySQL command-line client executing against the local Docker database. Each capture includes the SQL statement and its actual output; the cursor uses vertical result display. Figure 6.1 shows the tail of the actual integration-test log in a terminal. The capture script, SQL files, raw outputs, browser booking record and test log are retained under tmp/report for reproduction.')
p('Tests and the rollback demonstration leave no disposable records. The browser workflow retains one cancelled synthetic booking: the visit count is therefore 6,196 rather than 6,195. The phone-sized image is a responsive browser view; native hotspot acceptance remains pending.')

# Ensure body-font consistency and update fields on open; Word will refresh them before PDF export.
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
doc.core_properties.title=TITLE;doc.core_properties.subject='BCSE302P Database Systems Laboratory Mini Project'
doc.core_properties.author='Kartikay Gupta; Sreyanko Sinha';doc.core_properties.keywords='QureFlow, MySQL, BCSE302P, queue management'
doc.core_properties.comments=''
OUT.mkdir(parents=True,exist_ok=True)
path=OUT/'BCSE302P_Project_QureFlow.docx'
doc.save(path)
print(json.dumps({'docx':str(path),'abstractWords':len(abstract.split()),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'paragraphs':len(doc.paragraphs)},indent=2))
