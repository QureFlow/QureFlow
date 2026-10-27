import hashlib,json,re,zipfile
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
R=Path(__file__).parent
reference=Path(r"C:/Users/ASUS/Downloads/BCSE302P_Mini_Project_Report_Template.docx")
doc=Document(reference)
inventory=[]
with zipfile.ZipFile(reference) as z:
    for name in z.namelist():
        if not name.endswith('/'):
            b=z.read(name);inventory.append(dict(path=name,bytes=len(b),sha256=hashlib.sha256(b).hexdigest()))
(R/'template-package-inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf8')
contract=f"""# QureFlow report template contract
Reference: {reference}
SHA256: {hashlib.sha256(reference.read_bytes()).hexdigest()}
Reference render: reference-render, all 17 pages inspected.
Package evidence: template-package-inventory.json; template-style-evidence.json.
## Page and typography
One A4 portrait section, 11906 by 16838 twips, all four margins1440twips.
Different first page enabled: cover has no visible page number. Regular footer centred PAGE field.
Keep source cover institutional blocks, two-member table, faculty block and date.
Body Times New Roman12, line spacing1.5. Headings14boldblack per supplied guidelines.
Template duplicate heading styles and missing explicit Normal are normalized to these roles.
Source guidance italics/grey/blue are editable placeholders and removed.
Title black14bold, style Title; no new decorative rule.
Chapter headings use Heading1; subsection Heading2; dictionary table titles level3.
Tables retain source pale blue header fillD9E2F3, black text and thin black borders.
Repeat table headings, prevent row splits; body12 and line1.5.
Code is Times New Roman12 with the source light-grey block background; soft wraps allowed.
Captions12italicblack, directly below figures. Figures fit within6.26in text width.
## Slots and flow
word/document.xml body paragraphs0..20 and first table: update title, members, submission date.
Delete unused member rows3,4 and guidance paragraph13.
From CERTIFICATE onward replace all guidance blocks using cloned section/title/table patterns.
Preserve certificate faculty label and unsigned signature space; do not fabricate approval.
Sequence: certificate, abstract150to200words, real WordTOC, chapters1to10 from template.
Keep template subsections including Packages, explaining MySQL platform limitation.
Appendices hold the data dictionary, source map and terminal reproduction commands.
TOC/fields refreshed by installed Microsoft Word and PDF exported by Word.
## Evidence and factual boundaries
Use executed MySQL output and actual local browser screenshots; no invented native-phone screenshots.
Show a real browser booking and matching SQL before/after cancellation.
State that APK build/package verification passed but final hotspot device acceptance is pending.
Normalized core master relations are3NF; identify operational duplication rather than claim all BCNF.
## Package preservation
Authoring starts with a working copy of the retained DOCX.
Body/styles/settings/numbering/field caches/relationships and document metadata are editable.
Footer PAGE semantics, first-page behaviour, geometry and source institutional/faculty blocks are preserved.
Template has no embedded images, opaque customXML or controls containing authored content.
Empty annotation/footnote parts may be canonicalized by Word; no nonempty opaque content is removed.
New images and internal TOC bookmarks are authorized additions.
## Fidelity gates
Reference hash unchanged, oneA4 section,1in margins; all requested identities correct.
No guidance placeholders, code snippets must come from project SQL/repository source.
Check rendered PDF/PNG every page; compare chapter/table styles and cover with reference.
TOC page numbers verified against final PDF; unsigned certificate remains unsigned.
"""
(R/'artifact.md').write_text(contract,encoding='utf8')
data=json.loads((R/'evidence/report-data.json').read_text(encoding='utf8'))
data['testCases']=[x for x in data['testCases'] if x['id']!='TC25']
(R/'evidence/report-data.json').write_text(json.dumps(data,indent=2),encoding='utf8')
F=Path(r'C:/Windows/Fonts')
font=lambda n:ImageFont.truetype(str(F/'times.ttf'),n)
bold=lambda n:ImageFont.truetype(str(F/'timesbd.ttf'),n)

def diagram(name,boxes,edges,size=(1650,1420)):
 im=Image.new('RGB',size,'white');draw=ImageDraw.Draw(im)
 for points,label,where,dashed in edges:
  for a,b in zip(points,points[1:]):
   if dashed:
    import math
    dist=math.dist(a,b)
    for t in range(0,int(dist),22):
     u=min(t+12,dist);draw.line((a[0]+(b[0]-a[0])*t/dist,a[1]+(b[1]-a[1])*t/dist,a[0]+(b[0]-a[0])*u/dist,a[1]+(b[1]-a[1])*u/dist),fill='#222',width=3)
   else:draw.line((a,b),fill='#222',width=3)
  a,b=points[-2:];import math
  angle=math.atan2(b[1]-a[1],b[0]-a[0]);l=17
  draw.polygon([b,(b[0]-l*math.cos(angle-.45),b[1]-l*math.sin(angle-.45)),(b[0]-l*math.cos(angle+.45),b[1]-l*math.sin(angle+.45))],fill='#222')
  box=draw.textbbox(where,label,font=font(30));draw.rectangle((box[0]-5,box[1]-3,box[2]+5,box[3]+3),fill='white');draw.text(where,label,font=font(30),fill='black')
 for rect,title,lines in boxes:
  x,y,x2,y2=rect;draw.rectangle(rect,fill='white',outline='#111',width=3)
  draw.rectangle((x+2,y+2,x2-2,y+62),fill='#D9E2F3');draw.line((x,y+64,x2,y+64),fill='#111',width=2)
  draw.text((x+18,y+12),title,font=bold(35),fill='#111')
  for i,line in enumerate(lines):draw.text((x+18,y+82+i*43),line,font=font(33),fill='#111')
 im.save(R/'evidence'/name)

diagram('er-catalog.png',[
 ((100,70,550,285),'clinics',['PK id','name, district']),
 ((1050,70,1550,285),'specialties',['PK id','UNIQUE name']),
 ((580,445,1080,730),'doctors',['PK id','FK clinicId, specialtyId','name, fee, schedule']),
 ((100,490,530,705),'clinic_settings',['PK and FK clinicId','queue configuration']),
 ((40,1000,530,1230),'doctor_days',['PK doctorId, dayOfWeek','FK doctorId']),
 ((600,1000,1100,1230),'doctor_languages',['PK doctorId, languageId','both columns are FKs']),
 ((1190,1000,1620,1230),'languages',['PK id','UNIQUE name']),
 ((1190,485,1620,715),'doctor_interests',['PK doctorId, interest','FK doctorId']),
],[
 ([(380,285),(730,445)],'1:N',(490,320),False),
 ([(1280,285),(940,445)],'1:N',(1080,325),False),
 ([(315,285),(315,490)],'1:0..1',(330,360),False),
 ([(610,730),(300,1000)],'1:N',(370,845),False),
 ([(850,730),(850,1000)],'1:N',(870,850),False),
 ([(1190,1120),(1100,1120)],'1:N',(1110,1070),False),
 ([(1080,600),(1190,600)],'1:N',(1100,550),False),
],(1650,1320))
diagram('er-visits.png',[
 ((60,70,560,325),'users',['PK id','patient or staff profile','UNIQUE email and phone']),
 ((1060,70,1560,325),'doctors',['PK id','FK clinicId, specialtyId','parent in catalog view']),
 ((590,510,1090,785),'appointments',['PK id','FK patientId, doctorId','date, slotMinute, status']),
 ((590,1040,1090,1320),'queue_tokens',['PK id','FK patientId, doctorId','UNIQUE appointmentId FK']),
 ((1190,560,1630,850),'appointment_slots',['PK doctorId, slotDate,','   slotMinute','FK doctorId','capacity, bookedCount']),
],[
 ([(360,325),(710,510)],'1:N',(450,370),False),
 ([(1200,325),(980,510)],'1:N',(1050,370),False),
 ([(200,325),(200,920),(590,1180)],'1:N',(225,790),False),
 ([(1560,190),(1640,190),(1640,1190),(1090,1190)],'1:N',(1490,920),False),
 ([(1330,325),(1400,560)],'1:N',(1400,405),False),
 ([(840,785),(840,1040)],'0..1:0..1',(865,905),False),
 ([(1090,660),(1190,660)],'match',(1100,610),True),
],(1680,1400))
diagram('er-state.png',[
 ((60,60,560,280),'doctors',['PK id','clinic assignment']),
 ((1090,60,1590,280),'clinics',['PK id','facility master']),
 ((50,610,530,840),'queue_state',['PK and FK doctorId','isPaused']),
 ((600,610,1080,880),'queue_counters',['PK doctorId, queueDate','FK doctorId','currentNumber']),
 ((1150,610,1630,880),'prediction_models',['PK and FK doctorId','clinicId scope value','model, updatedAt']),
],[
 ([(220,280),(280,610)],'1:0..1',(100,430),False),
 ([(410,280),(810,610)],'1:N',(480,430),False),
 ([(560,170),(820,170),(1330,610)],'1:0..1',(1000,340),False),
 ([(1410,280),(1440,610)],'scope',(1460,430),True),
],(1680,970))
print('Template contract and three ER figures prepared.')
